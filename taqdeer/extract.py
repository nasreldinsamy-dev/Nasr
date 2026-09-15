import openpyxl, json, re, math
from openpyxl.utils import column_index_from_string, get_column_letter

SRC={'A':'out/BOQ-RAZEEN-A_15-09-2026_-_REV_02.xlsx','B':'work/B.xlsx'}
# rate-alignment results: the file's cached cost cells are stale for these rows
FIXCOST={'A':{4:{30:('=140*1.05+85',[('material',147.0),('installation',85.0)])},
              8:{36:('=105+20+0+18.5',[('cable',105.0),('installation',20.0),('testing',0.0),('accessories',18.5)]),
                 231:('=4600+8*2100',[('recorder',4600.0),('hard disks 8 x 2100',16800.0)])}}}
TOTROW={'A':{1:25,2:42,3:25,4:33,5:30,6:33,7:135,8:239,9:143,10:21},
        'B':{1:25,2:42,3:27,4:33,5:30,6:29,7:135,8:239,9:143,10:21}}
DIVNAME={1:'EARTHWORKS',2:'CIVIL WORKS',3:'MASONRY',4:'ARCHITECTURAL WORKS',
         5:'THERMAL & MOISTURE PROTECTION',6:'DOORS & WINDOWS',7:'PLUMBING',
         8:'ELECTRICAL',9:'FIRE PROTECTION SYSTEM',10:'METAL WORKS / STEEL STRUCTURE'}

def clean(v): return '' if v is None else ' '.join(str(v).split())

def label_for(wsf, col, row):
    for hr in range(row-1, max(0,row-14), -1):
        v=wsf.cell(hr,col).value
        if isinstance(v,str) and v.strip() and not v.startswith('=') and len(v.strip())<=24:
            return v.strip()
    return ''

out={'markup':1.25,'buildings':{}}
for b,path in SRC.items():
    wbf=openpyxl.load_workbook(path); wbv=openpyxl.load_workbook(path, data_only=True)
    divs=[]
    for d in range(1,11):
        sn=f'Div.{d}.{b}'; wsf=wbf[sn]; wsv=wbv[sn]
        fix=FIXCOST.get(b,{}).get(d,{})
        items=[]; group=''
        for r in range(7, TOTROW[b][d]+1):
            desc=clean(wsf.cell(r,3).value)
            f=wsf.cell(r,6).value
            if not (isinstance(f,str) and f.startswith('=')):
                if desc and len(desc)<160 and r!=TOTROW[b][d]: group=desc
                continue
            g=wsf.cell(r,7).value
            m=re.fullmatch(r'=F(\d+)\*E(\d+)', str(g))
            qrow=int(m.group(2)) if m else r
            try: qty=float(wsv.cell(qrow,5).value)
            except (TypeError,ValueError): qty=0.0

            hf=wsf.cell(r,8).value
            cost=wsv.cell(r,8).value
            res=[]; mode='fixed'
            if r in fix:
                hf, pairs = fix[r]
                res=[{'label':l,'value':v} for l,v in pairs]
                cost=sum(v for _,v in pairs)
                mode='sum' if '*' not in hf.split('=')[1].replace('8*2100','') else 'formula'
                mode='sum' if r!=231 else 'formula'
            elif isinstance(hf,str) and hf.startswith('='):
                body=hf[1:]
                cols=[]
                for cl,rw in re.findall(r'\$?([A-Z]{1,2})\$?(\d+)', body):
                    ci=column_index_from_string(cl)
                    if ci>=9 and int(rw)==r and ci not in cols: cols.append(ci)
                for ci in cols:
                    v=wsv.cell(r,ci).value
                    res.append({'label':label_for(wsf,ci,r) or f'Item {get_column_letter(ci)}',
                                'value':round(float(v),4) if isinstance(v,(int,float)) else 0.0})
                stripped=re.sub(r'\$?[A-Z]{1,2}\$?\d+','X',body)
                mode='sum' if res and re.fullmatch(r'[X+\s()]+', stripped) else ('formula' if res else 'fixed')
            cost=round(float(cost),4) if isinstance(cost,(int,float)) else 0.0
            price=math.ceil(float(f'{cost*1.25:.15g}'))
            items.append({'row':r,'code':clean(wsf.cell(r,1).value) or clean(wsf.cell(r,2).value),
                          'desc':desc,'unit':clean(wsf.cell(r,4).value),'qty':qty,'cost':cost,
                          'price':price,'f':str(hf) if isinstance(hf,str) else '',
                          'mode':mode,'res':res,'group':group})
        divs.append({'no':d,'name':DIVNAME[d],'items':items})
    out['buildings'][b]=divs

json.dump(out, open('boqdata.json','w'), ensure_ascii=False, separators=(',',':'))
import os, math
def x15(v): return float(f'{v:.15g}')
bad=[]; modes={}; sumbad=[]
for b,dv in out['buildings'].items():
    for D in dv:
        for it in D['items']:
            modes[it['mode']]=modes.get(it['mode'],0)+1
            if math.ceil(x15(it['cost']*1.25))!=it['price']: bad.append((b,D['no'],it['row'],it['cost'],it['price']))
            if it['mode']=='sum' and abs(sum(r['value'] for r in it['res'])-it['cost'])>0.02: sumbad.append((b,D['no'],it['row']))
print('size',os.path.getsize('boqdata.json')//1024,'KB | modes',modes)
print('price != ceil(cost*1.25):', bad or 'none')
print('sum-mode rows where resources do not add up to cost:', sumbad or 'none')
for b in 'AB':
    print(b,'total', f"{sum(i['qty']*i['price'] for D in out['buildings'][b] for i in D['items']):,.0f}")